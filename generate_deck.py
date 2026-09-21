import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Industrial Logistics Precision)
    COLOR_BG = RGBColor(11, 21, 35)         # #0B1523 Dark Navy
    COLOR_CARD = RGBColor(19, 34, 56)       # #132238 Surface Navy
    COLOR_CARD_BORDER = RGBColor(38, 59, 90)# Border
    COLOR_PRIMARY = RGBColor(234, 88, 12)   # #EA580C Highway Orange
    COLOR_SECONDARY = RGBColor(6, 182, 212) # #06B6D4 Cyan
    COLOR_SUCCESS = RGBColor(16, 185, 129)  # #10B981 Emerald
    COLOR_TEXT_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8 Slate 400
    COLOR_TEXT_DIM = RGBColor(100, 116, 139)   # #64748B Slate 500
    COLOR_ACCENT_BG = RGBColor(30, 58, 95)  # Accent Container

    logo_path = 'assets/logoICON.jpeg' if os.path.exists('assets/logoICON.jpeg') else None

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="NAMMA LORRY — TECHNICAL IMPLEMENTATION"):
        # Top category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(10), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_PRIMARY

        # Main slide title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(10.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_WHITE

        # Top Right mini logo / badge
        if logo_path:
            try:
                slide.shapes.add_picture(logo_path, Inches(12.0), Inches(0.5), height=Inches(0.7))
            except Exception:
                pass

        # Subtle divider line
        divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.55), Inches(11.733), Inches(0.02))
        divider.fill.solid()
        divider.fill.fore_color.rgb = COLOR_CARD_BORDER
        divider.line.fill.background()

    def add_card(slide, left, top, width, height, title="", border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)

        if title:
            tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = COLOR_SECONDARY

        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Decorative background glows
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.4), Inches(7.5))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_PRIMARY
    accent_bar.line.fill.background()

    if logo_path:
        try:
            slide1.shapes.add_picture(logo_path, Inches(1.2), Inches(1.4), height=Inches(1.2))
        except Exception:
            pass

    # Title box
    tb = slide1.shapes.add_textbox(Inches(1.2), Inches(2.7), Inches(11), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "NAMMA LORRY"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY

    p1 = tf.add_paragraph()
    p1.text = "Driver GPS & Distance Tracking System"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Complete Technical Implementation, Offline-First Architecture & Multi-Model Telemetry Engine"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(14)

    p3 = tf.add_paragraph()
    p3.text = "Specification Standard: gps.docx & DESIGN.md | Version 1.0.0 Production Release | Zero-Cost Infrastructure"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_SUCCESS
    p3.space_before = Pt(28)

    # =========================================================================
    # SLIDE 2: Executive Summary & The Problem in Indian Logistics
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Executive Summary: The Freight Mileage Dilemma in India")

    problems = [
        ("1. Rural Network Dropouts", "Cellular coverage routinely collapses across ghat sections, remote national corridors, and industrial transport nagars. Cloud-dependent apps halt logging, causing lost mileage.", COLOR_PRIMARY),
        ("2. GPS Multipath Sensor Noise", "Under highway flyovers and near massive metallic lorry bodies, smartphone GPS reflects erratically, triggering 2km jumps in 5s and inflating distances by 15% to 40%.", COLOR_SECONDARY),
        ("3. Rigid Static Baselines", "Fixed corridor tables (e.g. Chennai-Bangalore = 350km) ignore outer ring road bypasses, loading hub detours, and diversions, sparking frequent driver-operator disputes.", COLOR_SUCCESS),
        ("4. High Commercial Tool Costs", "Commercial AIS-140 hardware and proprietary routing APIs (Google Maps) impose prohibitive recurring monthly charges that erode thin logistics operating margins.", RGBColor(245, 158, 11))
    ]

    col_w = 2.7
    gap = 0.31
    for i, (title, desc, color) in enumerate(problems):
        x = 0.8 + i * (col_w + gap)
        add_card(slide2, x, 1.8, col_w, 3.8, title=title, border_color=color)

        tb = slide2.shapes.add_textbox(Inches(x + 0.25), Inches(2.45), Inches(col_w - 0.5), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.3

    # Bottom summary card
    add_card(slide2, 0.8, 5.85, 11.733, 1.15, title="")
    tb = slide2.shapes.add_textbox(Inches(1.05), Inches(5.95), Inches(11.2), Inches(0.95))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "THE SOLUTION: Zero Lost Mileage + Multi-Stage Quality Guard + Zero Infrastructure Cost"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS

    p2 = tf.add_paragraph()
    p2.text = "Namma Lorry introduces an offline-first telemetry platform on commodity Android phones using satellite GNSS, mathematical noise rejection, and multi-model arbitration across OpenStreetMap/OSRM and Supabase."
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_TEXT_WHITE
    p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 3: System Architecture & Workflow Pipeline
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "System Architecture: End-to-End Telemetry Pipeline")

    steps = [
        ("Step 1: GNSS Satellite Fix", "Hardware GPS captures lat, lng, speed, accuracy, and timestamp without needing internet.", COLOR_SECONDARY),
        ("Step 2: 3-Stage Quality Guard", "Real-time mathematical filter rejects accuracy >50m, speeds >120 km/h, and multipath spikes.", COLOR_PRIMARY),
        ("Step 3: Offline Local Queue", "Points stamped with sync_status: pending and stored locally during network blackouts.", COLOR_SUCCESS),
        ("Step 4: Supabase Auto-Sync", "Upon reconnection, points batch-sync via idempotent upsert with zero duplicate records.", COLOR_SECONDARY),
        ("Step 5: Multi-Model Arbiter", "Evaluates Model A (Baseline), Model B (GPS Haversine), and Model C (OSRM Road Snapped).", COLOR_PRIMARY),
        ("Step 6: Driver HUD & Fleet Command", "Renders live speedometer, trip audit breakdown, daily analytics, and admin map.", COLOR_SUCCESS)
    ]

    for i, (title, desc, color) in enumerate(steps):
        row = i // 3
        col = i % 3
        x = 0.8 + col * (3.7 + 0.31)
        y = 1.8 + row * (2.55 + 0.25)

        add_card(slide3, x, y, 3.7, 2.55, title=title, border_color=color)

        tb = slide3.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.65), Inches(3.2), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.3

    # =========================================================================
    # SLIDE 4: Zero-Cost Open-Source Technology Stack
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Zero-Cost Open-Source Technology Stack")

    stack_items = [
        ("Frontend Application", "Vanilla ES6+ & HTML5 Semantic UI", "Zero bundle overhead, ultra-fast rendering on budget Android devices", "100% Free"),
        ("Design System", "Vanilla CSS3 Tokens (Industrial Logistics)", "8dp grid, tabular numerals, Material Design 3 elevation and contrast", "100% Free"),
        ("Map Display Engine", "Leaflet.js v1.9.4", "Touch-optimized, lightweight interactive mapping with custom lorry icons", "Open-Source"),
        ("Map Tile Cartography", "OpenStreetMap (OSM) Standard Tiles", "Global street-level highway and corridor cartographic tile server", "Open-Source"),
        ("Road Routing Engine", "OSRM (Open Source Routing Machine)", "Topological road-network matching, highway geometry, turn-by-turn routing", "Open-Source"),
        ("Cloud Database", "Supabase (PostgreSQL 15+)", "Relational store for drivers, lorries, trips, and GPS location history", "Free Tier (500MB)"),
        ("Spatial Geodatabase", "PostGIS Extension", "Native spatial indexing, geometric coordinates, and distance projections", "Built-in Free"),
        ("Cloud Hosting & CDN", "Vercel Serverless Edge Platform", "Global edge delivery, automatic HTTPS SSL, zero maintenance deployment", "Free Tier")
    ]

    # Create 2 columns of 4 cards
    card_w = 5.7
    card_h = 1.15
    for i, (layer, tech, detail, cost) in enumerate(stack_items):
        col = i // 4
        row = i % 4
        x = 0.8 + col * (card_w + 0.33)
        y = 1.8 + row * (card_h + 0.16)

        add_card(slide4, x, y, card_w, card_h, title="")

        tb = slide4.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.12), Inches(card_w - 0.5), Inches(card_h - 0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"{layer}: {tech}"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_TEXT_WHITE

        p1 = tf.add_paragraph()
        p1.text = f"{detail} — [{cost}]"
        p1.font.size = Pt(11)
        p1.font.color.rgb = COLOR_SECONDARY
        p1.space_before = Pt(3)

    # =========================================================================
    # SLIDE 5: Database Architecture & PostGIS Data Models
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Database Architecture & PostGIS Integration")

    tables = [
        ("drivers (Section 4.1)", "• id (TEXT PRIMARY KEY)\n• name, phone, email\n• status ('active' | 'idle')\n• created_at (TIMESTAMPTZ)\n\nStores verified lorry drivers with phone authentication."),
        ("vehicles (Section 5)", "• id (TEXT PRIMARY KEY)\n• driver_id (FK -> drivers.id)\n• vehicle_number (e.g. TN 01 AB 1234)\n• vehicle_type ('25 Ton Heavy Truck')\n• capacity, status ('on_road')\n\nDirect driver-to-lorry asset binding."),
        ("trips (Section 6)", "• id (TEXT PRIMARY KEY)\n• driver_id, vehicle_id (FKs)\n• start_location, end_location\n• started_at, ended_at, status\n• distance_km (NUMERIC 10,2)\n• calculation_method, audit_confidence\n\nCanonical audited journey ledger."),
        ("driver_locations (Section 7)", "• id (TEXT PRIMARY KEY)\n• trip_id, driver_id (FKs)\n• latitude, longitude (DOUBLE)\n• accuracy, speed, heading\n• recorded_at, synced_at\n\nComposite Index: (trip_id, recorded_at DESC)\nHigh-frequency GNSS coordinate logs.")
    ]

    card_w = 2.7
    gap = 0.31
    for i, (title, content) in enumerate(tables):
        x = 0.8 + i * (card_w + gap)
        add_card(slide5, x, 1.8, card_w, 4.0, title=title, border_color=COLOR_SECONDARY)

        tb = slide5.shapes.add_textbox(Inches(x + 0.25), Inches(2.45), Inches(card_w - 0.5), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.3

    # Bottom RLS bar
    add_card(slide5, 0.8, 6.0, 11.733, 1.0, title="", border_color=COLOR_SUCCESS)
    tb = slide5.shapes.add_textbox(Inches(1.05), Inches(6.1), Inches(11.2), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ROW LEVEL SECURITY (RLS) & PRIVACY (Section 28)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p1 = tf.add_paragraph()
    p1.text = "Granular PostgreSQL policies guarantee Driver A cannot read Driver B's telemetry or trips. Fleet managers access authorized fleets under organizational scopes."
    p1.font.size = Pt(11)
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_before = Pt(3)

    # =========================================================================
    # SLIDE 6: Hardware GPS Acquisition & Calibration
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Hardware GPS Telemetry & Pre-Trip Calibration")

    cards_s6 = [
        ("1. HTML5 Geolocation Watcher", "Direct integration with device hardware via navigator.geolocation.watchPosition(). Configured with enableHighAccuracy: true, maximumAge: 2000ms, and timeout: 12000ms for continuous streaming.", COLOR_SECONDARY),
        ("2. Pre-Trip GNSS Locking", "Acquires an initial satellite fix before departure. Drivers can press 'Lock / Re-acquire Start GPS' to pin exact depot bay coordinates and ensure zero initial displacement errors.", COLOR_PRIMARY),
        ("3. Reverse Geocode Resolution", "Coordinates are queried asynchronously against OpenStreetMap Nominatim to resolve human-readable transport addresses (e.g. 'Madhavaram Truck Terminal, Chennai').", COLOR_SUCCESS),
        ("4. Raw Telemetry Attributes", "Every recorded point captures: Latitude (WGS-84), Longitude, Accuracy (m), Speed (m/s converted to km/h via * 3.6), Heading (0°-360°), and Hardware Timestamp.", RGBColor(245, 158, 11))
    ]

    card_w = 5.7
    card_h = 2.45
    for i, (title, content, color) in enumerate(cards_s6):
        col = i % 2
        row = i // 2
        x = 0.8 + col * (card_w + 0.33)
        y = 1.8 + row * (card_h + 0.25)

        add_card(slide6, x, y, card_w, card_h, title=title, border_color=color)

        tb = slide6.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.65), Inches(card_w - 0.5), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.35

    # =========================================================================
    # SLIDE 7: 3-Stage GPS Quality Guard & Filtering Engine
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "3-Stage GPS Quality Guard & Anomaly Filtering Engine")

    checks = [
        ("Check 1: Accuracy Filter", "Rule: accuracy <= 50.0 meters\n\n• Satellite triangulation errors frequently range between 300m - 1500m during cell handovers.\n• Any point reporting an error radius > 50m is flagged POOR_ACCURACY and dropped immediately.\n• Prevents massive geometric inflation.", COLOR_SECONDARY),
        ("Check 2: Speed Jump Filter", "Rule: v = (Δd / Δt) * 3.6 <= 120 km/h\n\n• Computes instantaneous velocity between consecutive valid points.\n• Heavy freight lorries in India cannot exceed 120 km/h.\n• Points exceeding 120 km/h with displacement > 500m are rejected as IMPOSSIBLE_SPEED anomalies.", COLOR_PRIMARY),
        ("Check 3: 3-Point Spike Filter", "Rule: Detour > 2.2x and (Detour - Direct) > 400m\n\n• Evaluates 3-point sliding windows (P_i-1 -> P_i -> P_i+1).\n• Flyovers and metallic cargo cause transient single-point reflections.\n• If P_i creates an acute deflection > 400m, it is dropped as TRANSIENT_SPIKE.", COLOR_SUCCESS)
    ]

    card_w = 3.7
    gap = 0.31
    for i, (title, content, color) in enumerate(checks):
        x = 0.8 + i * (card_w + gap)
        add_card(slide7, x, 1.8, card_w, 3.8, title=title, border_color=color)

        tb = slide7.shapes.add_textbox(Inches(x + 0.25), Inches(2.45), Inches(card_w - 0.5), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.3

    # Dynamic interval box
    add_card(slide7, 0.8, 5.85, 11.733, 1.15, title="")
    tb = slide7.shapes.add_textbox(Inches(1.05), Inches(5.95), Inches(11.2), Inches(0.95))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "BATTERY OPTIMIZATION & ADAPTIVE SAMPLING INTERVALS (Section 10)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY

    p2 = tf.add_paragraph()
    p2.text = "• Stationary / Traffic Jam (v < 5 km/h): Capture every 30s | City Driving (5-50 km/h): Capture every 10s | Highway Cruising (v > 50 km/h): Capture every 5s. Conserves phone battery across 12-hour hauls."
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_TEXT_WHITE
    p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 8: Multi-Model Distance Engine (Models A, B, C)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Multi-Model Distance Calculation Engine (A / B / C)")

    models = [
        ("MODEL A — Known Baseline & Fallback", "• Standard Freight Corridor Baseline:\n  - Chennai <-> Bengaluru: 348.5 KM\n  - Chennai <-> Sriperumbudur: 42.0 KM\n  - Chennai <-> Vellore: 138.0 KM\n  - Ranipet <-> Hosur: 235.6 KM\n\n• Dynamic Straight Line * 1.25 Winding Factor:\n  For uncatalogued routes, multiplies straight-line displacement by 1.25x (National Highway curvature).\n• Strict Role: Emergency fallback only.", COLOR_PRIMARY),
        ("MODEL B — GPS Geodesic Haversine", "• 100% Offline-Capable Calculation:\n  Accumulates great-circle distances between consecutive validated points:\n  d = Σ Haversine(P_i, P_i+1)\n\n• Earth radius R = 6,371 km.\n• Operates entirely on the smartphone without cellular internet.\n• Reflects the driver's actual physical path, including loading detours and bypasses.", COLOR_SUCCESS),
        ("MODEL C — OSM + OSRM Road Matching", "• Snapped Road Network Distance:\n  Queries OpenStreetMap OSRM routing engine with verified start and stop coordinates.\n\n• High-Fidelity Geometry:\n  Snaps vehicle coordinates to real road network topology and highway lane geometry.\n• Direct Start-to-End Routing prevents intermediate waypoint U-turn penalties across highway medians.", COLOR_SECONDARY)
    ]

    card_w = 3.7
    gap = 0.31
    for i, (title, content, color) in enumerate(models):
        x = 0.8 + i * (card_w + gap)
        add_card(slide8, x, 1.8, card_w, 5.2, title=title, border_color=color)

        tb = slide8.shapes.add_textbox(Inches(x + 0.25), Inches(2.55), Inches(card_w - 0.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.35

    # =========================================================================
    # SLIDE 9: Algorithmic Arbiter & Decision Reconciliation
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Algorithmic Arbiter & Distance Reconciliation Matrix")

    rules = [
        ("1. Stationary Vehicle Lock (0.0 KM)", "If start and stop coordinates show displacement < 60 meters and total movement < 80 meters, the trip is verified at 0.0 KM. Prevents drivers from billing for stationary yard idling.", COLOR_SUCCESS),
        ("2. Physical Boundary Verification", "Both Model B (GPS) and Model C (OSRM) must be >= 0.85 * straight-line displacement. A road trajectory cannot physically be shorter than straight line.", COLOR_SECONDARY),
        ("3. Outlier Detour Rejection", "If OSRM reports extreme detours (Model C > 2.8 * displacement), it is rejected as a median turnaround artifact in favor of verified driver GPS trajectory.", COLOR_PRIMARY),
        ("4. Multi-Model Priority Hierarchy", "If B & C agree within 35% -> Model B (Road Confirmed).\nIf GPS corrupted -> Model C (OSRM Road Network).\nIf OSRM offline -> Model B (GPS Filtered).\nIf both fail -> Model A (Fallback) or NEEDS_REVIEW.", RGBColor(245, 158, 11))
    ]

    card_w = 5.7
    card_h = 2.45
    for i, (title, content, color) in enumerate(rules):
        col = i % 2
        row = i // 2
        x = 0.8 + col * (card_w + 0.33)
        y = 1.8 + row * (card_h + 0.25)

        add_card(slide9, x, y, card_w, card_h, title=title, border_color=color)

        tb = slide9.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.65), Inches(card_w - 0.5), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.35

    # =========================================================================
    # SLIDE 10: Offline-First Resilience & Sync Engine
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Offline-First Resilience & Zero-Duplicate Sync Engine")

    flow_items = [
        ("Offline Point Capture", "When cellular internet drops (isOnline = false), GPS logging continues. Points are stamped with sync_status: 'pending' and safely appended to local storage (IndexedDB / LocalStorage).", COLOR_PRIMARY),
        ("Realtime Network Listener", "Engine attaches listeners to browser 'online' and 'offline' events. Includes manual simulation toggle in header for QA and driver testing.", COLOR_SECONDARY),
        ("Atomic Batch Synchronization", "Upon reconnection, pending queue points are batched into chunks of 50. Mutex lock ensures no concurrent sync races or duplicate request threads.", COLOR_SUCCESS),
        ("Zero-Duplicate Upsert Deduplication", "Supabase receives points via: upsert({ ...point }, { onConflict: 'id', ignoreDuplicates: true }). Guaranteed zero duplicate records in cloud database.", RGBColor(245, 158, 11))
    ]

    card_w = 5.7
    card_h = 2.45
    for i, (title, content, color) in enumerate(flow_items):
        col = i % 2
        row = i // 2
        x = 0.8 + col * (card_w + 0.33)
        y = 1.8 + row * (card_h + 0.25)

        add_card(slide10, x, y, card_w, card_h, title=title, border_color=color)

        tb = slide10.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.65), Inches(card_w - 0.5), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.35

    # =========================================================================
    # SLIDE 11: Driver Mobile Cockpit & User Interface
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "Driver Mobile Cockpit HUD & Visual Telemetry")

    ui_elements = [
        ("Live HUD Telemetry Tiles", "• Speedometer (km/h) with tabular numerals to eliminate horizontal display jitter.\n• Trip Distance (KM) updating continuously in Highway Orange.\n• Trip Duration Clock (hh:mm:ss).\n• GNSS Accuracy Indicator (±m precision circle).", COLOR_PRIMARY),
        ("Dynamic Origin & Stop Cards", "• Real Start Location: Detects driver's exact starting point and street address.\n• Stopping Point: Captured immediately when driver taps 'STOP TRIP'.\n• Visual connection stepper with animated status indicators.", COLOR_SECONDARY),
        ("Interactive Leaflet Vector Map", "• Custom Animated Lorry Marker: Rotates heading dynamically based on vehicle bearing.\n• Green Polyline: Verified GPS trajectory.\n• Cyan Polyline: OSRM highway road snapped geometry.\n• Red Markers: Filtered noise anomalies.", COLOR_SUCCESS),
        ("Supervisor Trip Audit Modal", "• Opens automatically on trip completion.\n• Full multi-model comparison table (Model A vs B vs C).\n• Data Quality Breakdown: Raw points vs noise rejected.\n• Algorithmic explanation justifying the selected distance.", RGBColor(245, 158, 11))
    ]

    card_w = 5.7
    card_h = 2.45
    for i, (title, content, color) in enumerate(ui_elements):
        col = i % 2
        row = i // 2
        x = 0.8 + col * (card_w + 0.33)
        y = 1.8 + row * (card_h + 0.25)

        add_card(slide11, x, y, card_w, card_h, title=title, border_color=color)

        tb = slide11.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.65), Inches(card_w - 0.5), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.35

    # =========================================================================
    # SLIDE 12: Driver Analytics & Admin Fleet Command
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)
    add_header(slide12, "Driver Analytics & Admin Fleet Command")

    analytics_cards = [
        ("Driver KPI Dashboard", "• Today's Total Distance (KM)\n• Trips Completed Today\n• Total Active Duty Driving Hours\n• Current Month Accumulated KM\n\nDerived exclusively from genuine completed trips with zero dummy data padding.", COLOR_SECONDARY),
        ("Weekly Mon-Sun Distance Chart", "• Interactive 7-day bar chart visualization.\n• Displays daily kilometer distribution from Monday to Sunday.\n• Dynamic bar fill height calculations with responsive hover tooltips.\n• Instant weekly total distance summary.", COLOR_PRIMARY),
        ("Admin Fleet Map Surveillance", "• Live multi-lorry tracking across Tamil Nadu & Karnataka freight corridors.\n• Real-time vehicle status indicators (Available, On Road, Maintenance).\n• Displays lorry payload capacity (e.g. 25 Ton Heavy Truck, 32 Ton Multi-Axle).\n• Instant driver profile assignment.", COLOR_SUCCESS),
        ("Supabase Realtime CDC (Section 27)", "• Enables Supabase PostgreSQL Realtime WebSocket stream.\n• Dispatches live lorry coordinates to Fleet Command map without manual browser refresh.\n• Instant operational visibility for transport supervisors.", RGBColor(245, 158, 11))
    ]

    card_w = 5.7
    card_h = 2.45
    for i, (title, content, color) in enumerate(analytics_cards):
        col = i % 2
        row = i // 2
        x = 0.8 + col * (card_w + 0.33)
        y = 1.8 + row * (card_h + 0.25)

        add_card(slide12, x, y, card_w, card_h, title=title, border_color=color)

        tb = slide12.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.65), Inches(card_w - 0.5), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.35

    # =========================================================================
    # SLIDE 13: Automated Test Suite (Section 31 Testing Plan)
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13)
    add_header(slide13, "Automated Verification Suite: Section 31 Testing Plan")

    tests = [
        ("Test 1: Internet ON", "Verifies live GPS ingestion, Model B Haversine, OSRM road matching, and immediate cloud sync.", "PASSED", COLOR_SUCCESS),
        ("Test 2: Internet OFF", "Verifies network disconnection, continuous tracking, and local queue storage with status pending.", "PASSED", COLOR_SUCCESS),
        ("Test 3: Reconnection Sync", "Verifies pending points flush to Supabase, offline queue clears to 0, and 0 duplicate IDs exist.", "PASSED", COLOR_SUCCESS),
        ("Test 4: GPS Jump Simulation", "Injects 85m accuracy error and 1584 km/h jump. Confirms Check 1 & Check 2 block all anomalies.", "PASSED", COLOR_SUCCESS),
        ("Test 5: Adaptive Interval", "Validates dynamic interval scaling across stationary (30s), city (10s), and highway speeds (5s).", "PASSED", COLOR_SUCCESS)
    ]

    for i, (title, desc, status, color) in enumerate(tests):
        y = 1.8 + i * (0.95 + 0.12)
        add_card(slide13, 0.8, y, 11.733, 0.95, title="")

        tb = slide13.shapes.add_textbox(Inches(1.05), Inches(y + 0.12), Inches(9.5), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_TEXT_WHITE

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(11)
        p1.font.color.rgb = COLOR_TEXT_MUTED
        p1.space_before = Pt(2)

        # Status badge
        badge = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.8), Inches(y + 0.22), Inches(1.4), Inches(0.48))
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_ACCENT_BG
        badge.line.color.rgb = color
        badge.line.width = Pt(1.5)

        tb_b = slide13.shapes.add_textbox(Inches(10.8), Inches(y + 0.26), Inches(1.4), Inches(0.4))
        tf_b = tb_b.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = f"✓ {status}"
        p_b.font.size = Pt(12)
        p_b.font.bold = True
        p_b.font.color.rgb = color

    # =========================================================================
    # SLIDE 14: Deployment, Security & Production Roadmap
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14)
    add_header(slide14, "Deployment, Security & Production Scaling Roadmap")

    scaling_stages = [
        ("STAGE 1: Zero-Cost MVP", "• Fleet: 1 - 100 Trucks\n• DB: Supabase Free Tier (500MB DB, 50k MAU)\n• Routing: Public OSRM API\n• Hosting: Vercel Free Edge CDN\n• Total Cost: ₹0 / month", COLOR_SUCCESS),
        ("STAGE 2: Regional Fleet", "• Fleet: 100 - 1,000 Trucks\n• DB: Supabase Pro ($25/mo) + Monthly PostGIS Table Partitioning\n• Routing: Self-hosted OSRM Docker ($15/mo)\n• Total Cost: ~$40 / month", COLOR_SECONDARY),
        ("STAGE 3: National Enterprise", "• Fleet: 1,000 - 10,000+ Trucks\n• DB: Managed PostgreSQL / TimescaleDB\n• Routing: High-Availability OSRM Cluster\n• Total Cost: ~$250 / month", COLOR_PRIMARY)
    ]

    card_w = 3.7
    gap = 0.31
    for i, (title, content, color) in enumerate(scaling_stages):
        x = 0.8 + i * (card_w + gap)
        add_card(slide14, x, 1.8, card_w, 3.4, title=title, border_color=color)

        tb = slide14.shapes.add_textbox(Inches(x + 0.25), Inches(2.45), Inches(card_w - 0.5), Inches(2.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.line_spacing = 1.35

    # Bottom roadmap box
    add_card(slide14, 0.8, 5.45, 11.733, 1.55, title="IMMEDIATE PRODUCTION ROADMAP & INTEGRATIONS", border_color=RGBColor(245, 158, 11))
    tb = slide14.shapes.add_textbox(Inches(1.05), Inches(5.95), Inches(11.2), Inches(0.95))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "1. PWA Service Worker: Complete offline asset caching for instant launch in zero-connectivity areas.\n2. Background Geolocation Native Plugin: Migration to Capacitor/Flutter for tracking while screen is locked.\n3. E-Way Bill & FASTag Cross-Referencing: Automatic toll plaza timestamp verification against driver GPS trace.\n4. Driver Safety Scoring: Tracking sudden acceleration, harsh braking, and idling hours."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_WHITE
    p.line_spacing = 1.3

    # Save presentation
    output_path = "Namma_Lorry_GPS_Implementation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] Successfully created {output_path} with 14 widescreen slides.")

if __name__ == "__main__":
    build_presentation()

